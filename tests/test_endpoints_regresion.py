from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient

from app.api import compras as compras_api
from app.api.main import app
from app.core import concurrencia
from app.core.deps import get_current_user
from app.db.database import get_db

USUARIO_ACTUAL = SimpleNamespace(id=1, rol="admin", nombre="Test")


class FakeDB:
    """Solo responde a db.get(LoteProcesamiento, id): el lote 99 pertenece al usuario 2."""

    def get(self, _modelo, lote_id):
        return SimpleNamespace(id=lote_id, usuario_id=2) if lote_id == 99 else None


@pytest.fixture
def cliente(monkeypatch):
    monkeypatch.setattr(concurrencia, "_en_vuelo", 0)
    app.dependency_overrides[get_current_user] = lambda: USUARIO_ACTUAL
    app.dependency_overrides[get_db] = lambda: FakeDB()
    yield TestClient(app)
    app.dependency_overrides.clear()


def _archivos_minimos():
    return {
        "rg_files": ("rg.xlsx", b"esto no es un excel", "application/octet-stream"),
        "pos_data_json": ("pos.json", b"[]", "application/json"),
    }


def test_no_se_puede_reusar_un_lote_de_otro_usuario(cliente):
    # Regresión del IDOR corregido: lote_id 99 pertenece al usuario 2, no al 1.
    r = cliente.post("/api/compras/reconcile", files=_archivos_minimos(), data={"lote_id": "99"})
    assert r.status_code == 404


def test_el_libro_debe_llegar_como_archivo_y_no_como_texto(cliente):
    # Regresión del guard de Blob: un frontend viejo mandaba el JSON como texto plano.
    r = cliente.post(
        "/api/compras/reconcile",
        files={"rg_files": ("rg.xlsx", b"x", "application/octet-stream")},
        data={"pos_data_json": "[]"},
    )
    assert r.status_code == 422
    assert "texto plano" in r.json()["detail"]


def test_rg_corrupta_devuelve_422_con_mensaje_claro_no_500(cliente):
    # Regresión del fix de archivo corrupto en /api/compras/reconcile.
    r = cliente.post("/api/compras/reconcile", files=_archivos_minimos())
    assert r.status_code == 422
    assert r.json()["detail"] == compras_api.MSG_ARCHIVO_NO_VALIDO


def test_ingest_compras_con_duplicados_en_el_libro_devuelve_el_listado_en_el_paso_1(cliente, monkeypatch):
    from tests.fabricas import fila_compras as f

    filas = [f("001-001-000069580129373", 1.0), f("001-001-000069580129373", 2.0), f("001-001-000070080129373", 3.0)]
    monkeypatch.setattr(compras_api.engine, "ingest_file", lambda *a, **k: filas)
    r = cliente.post("/api/compras/ingest", files={"files": ("libro.xls", b"x", "application/octet-stream")})
    assert r.status_code == 422
    d = r.json()["detail"]
    assert d["tipo"] == "comprobantes_duplicados" and d["origen"] == "Libro"
    assert d["detalle"] == [{"comprobante": "001-001-0000695", "ruc": "80129373", "tipo": "Factura", "cantidad": 2, "origen": "Libro"}]


def test_excel_de_ventas_incluye_la_columna_fecha_junto_a_documento(cliente):
    from io import BytesIO
    from openpyxl import load_workbook

    lado = {"gravada_10": "0,00", "iva_10": "0,00", "gravada_5": "0,00", "iva_5": "0,00", "exenta": "0,00", "total": "0,00"}
    fila = {"doc": "001-001-0000001", "fecha": "05/10/2026", "tipo_doc": "Factura", "sistema": "S", "local": "L1",
            "libro": lado, "rg90": lado, "diferencia": "Coincide", "ruc": "80012345-6", "nombre": "Cliente Uno"}
    r = cliente.post("/api/export/diff-ventas-excel", json={"rows": [fila], "visible_columns": ["doc", "fecha", "ruc", "nombre", "tipo_doc"]})
    assert r.status_code == 200
    ws = load_workbook(BytesIO(r.content)).active
    filas = list(ws.iter_rows(values_only=True))
    assert filas[0][:5] == ("Documento", "Fecha", "RUC Cliente", "Nombre Cliente", "Tipo")
    assert filas[1][:5] == ("001-001-0000001", "05/10/2026", "80012345-6", "Cliente Uno", "Factura")
